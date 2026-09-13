"""Extract text from uploaded documents, with lazy optional OCR support."""

import json
import re
from pathlib import Path

_ocr_engine = None
_ocr_failed = False


def _normalize(value):
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", str(value).lower())


def _extract_pdf(path: Path):
    try:
        import pymupdf
    except ImportError:
        return "", "pdf-parser-unavailable", 0.0
    with pymupdf.open(path) as document:
        text = "\n".join(page.get_text("text") for page in document)
    if len(_normalize(text)) >= 20:
        return text, "pdf-text", 1.0
    return _extract_with_paddle(path)


def _extract_docx(path: Path):
    try:
        from docx import Document
    except ImportError:
        return "", "docx-parser-unavailable", 0.0
    document = Document(path)
    paragraphs = [item.text for item in document.paragraphs if item.text.strip()]
    for table in document.tables:
        paragraphs.extend(
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
        )
    text = "\n".join(paragraphs)
    return text, "docx-text", 1.0 if text.strip() else 0.0


def _load_ocr():
    global _ocr_engine, _ocr_failed
    if _ocr_engine is not None or _ocr_failed:
        return _ocr_engine
    try:
        from paddleocr import PaddleOCR

        _ocr_engine = PaddleOCR(
            lang="ch",
            use_doc_orientation_classify=True,
            use_doc_unwarping=True,
            use_textline_orientation=True,
        )
    except Exception:
        _ocr_failed = True
    return _ocr_engine


def _extract_with_paddle(path: Path):
    engine = _load_ocr()
    if engine is None:
        return "", "ocr-unavailable", 0.0
    texts, scores = [], []
    for result in engine.predict(input=str(path)):
        payload = getattr(result, "json", result)
        if callable(payload):
            payload = payload()
        if isinstance(payload, str):
            payload = json.loads(payload)
        data = payload.get("res", payload) if isinstance(payload, dict) else {}
        texts.extend(str(item) for item in data.get("rec_texts", []) if item)
        scores.extend(float(item) for item in data.get("rec_scores", []) if item is not None)
    confidence = sum(scores) / len(scores) if scores else (0.7 if texts else 0.0)
    return "\n".join(texts), "paddleocr", round(confidence, 3)


def extract_document(path: Path, mime_type: str):
    suffix = path.suffix.lower()
    if suffix == ".pdf" or mime_type == "application/pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in {".png", ".jpg", ".jpeg"} or mime_type.startswith("image/"):
        return _extract_with_paddle(path)
    if suffix == ".txt" or mime_type.startswith("text/"):
        return path.read_text(encoding="utf-8", errors="replace"), "plain-text", 1.0
    return "", "unsupported-format", 0.0
